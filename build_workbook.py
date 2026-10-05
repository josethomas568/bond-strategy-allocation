"""
Excel backup for the investor one-pager.

The workbook is live: adjusted ETF prices are the only data, and every return, portfolio,
statistic and stress figure is a formula. Changing the base weights, the allocation sizes
or a stress-window date on the Assumptions sheet recalculates everything.

Monthly rebalancing in formulas: for each holding, column "growth since month-end"
G_t = (1+r_t) on the first trading day of a month, else G_{t-1}*(1+r_t).
A portfolio rebalanced to weights w at each month-end then has daily return
  sum(w*G_t) / sum(w*G_{t-1}) - 1   (denominator = 1 on the first day of a month),
which is exactly a buy-and-drift portfolio reset to target weights every month.
"""
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.table import Table  # noqa: F401  (kept for readers extending the file)

from analysis import STRATEGIES, STRESS, EQUITY, BONDS, CASH, ROLL, START

FONT = "Arial"
BLUE, BLACK, GREEN = "0000FF", "000000", "008000"
HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
SUB_FILL = PatternFill("solid", fgColor="E8EDF3")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
THIN = Side(style="thin", color="C9D1DB")
PCT, PCT2, NUM2, DATE = "0.0%;-0.0%;-", "0.00%", "0.00", "yyyy-mm-dd"
PP = '+0.0" pts";-0.0" pts";0.0" pts"'

ASSETS = [EQUITY, BONDS, CASH, *STRATEGIES]                 # price / return column order
GROWTH = [EQUITY, BONDS, *STRATEGIES]                       # assets that get growth / peak / dd / vol columns


def f(bold=False, color=BLACK, size=10, italic=False):
    return Font(name=FONT, bold=bold, color=color, size=size, italic=italic)


def header(ws, row, labels, col=1, widths=None):
    for i, lab in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=lab)
        c.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 30
    if widths:
        for i, w in enumerate(widths):
            ws.column_dimensions[L(col + i)].width = w


def title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = f(True, size=14)
    if sub:
        ws["A2"] = sub
        ws["A2"].font = f(color="555555", italic=True)
    ws.sheet_view.showGridLines = False


def build(R, path):
    px = R["monthly_prices"]  # not used for formulas; daily prices below
    rets = R["daily_returns"]
    prices = R["prices_daily"]                       # includes the base day before START
    n = len(prices)                                  # rows 2..n+1 hold prices
    first, last = 3, n + 1                           # return rows in 'Daily calc'
    port_names = list(R["port_returns"].columns)

    wb = Workbook()
    ws_read = wb.active
    ws_read.title = "Read me"
    ws_ass = wb.create_sheet("Assumptions")
    ws_sum = wb.create_sheet("Summary")
    ws_alloc = wb.create_sheet("Allocation scenarios")
    ws_str = wb.create_sheet("Stress tests")
    ws_prof = wb.create_sheet("Strategy profiles")
    ws_cor = wb.create_sheet("Correlations")
    ws_px = wb.create_sheet("Prices")
    ws_d = wb.create_sheet("Daily calc")

    # ------------------------------------------------------------------ Assumptions
    title(ws_ass, "Assumptions", "Blue cells on yellow are inputs. Change them and the whole workbook recalculates.")
    ws_ass.column_dimensions["A"].width = 38
    for col, w in zip("BCDE", (14, 14, 14, 60)):
        ws_ass.column_dimensions[col].width = w
    rows = [
        (4, "Stock weight in the base 60/40 (SPY)", 0.60, PCT, "Typical balanced investor"),
        (5, "Bond weight in the base 60/40 (AGG)", 0.40, PCT, "Must add to 100% with the stock weight"),
        (8, "Strategy allocation, scenario 1", 0.10, PCT, "Funded pro rata from stocks and bonds"),
        (9, "Strategy allocation, scenario 2", 0.20, PCT, "Funded pro rata from stocks and bonds"),
        (17, "Trading days per year", 252, "0", "Used to annualize returns and volatility"),
    ]
    for r, lab, v, fmt, note in rows:
        ws_ass.cell(row=r, column=1, value=lab).font = f()
        c = ws_ass.cell(row=r, column=2, value=v)
        c.font, c.fill, c.number_format = f(color=BLUE), INPUT_FILL, fmt
        ws_ass.cell(row=r, column=5, value=note).font = f(color="555555", italic=True)
    ws_ass["A6"] = "Check: weights add to 100%"
    ws_ass["B6"] = '=IF(ABS(B4+B5-1)<0.000001,"OK","Weights must add to 100%")'
    ws_ass["A3"], ws_ass["A7"] = "Base portfolio", "Allocation sizes"
    ws_ass["A11"] = "Stress windows (return from the close on the start date to the close on the end date)"
    for c in ("A3", "A7", "A11", "A16"):
        ws_ass[c].font = f(True)
    header(ws_ass, 12, ["Episode", "Start", "End", "", "What happened"])
    for i, (k, (a, b, desc)) in enumerate(STRESS.items()):
        r = 13 + i
        ws_ass.cell(row=r, column=1, value=k).font = f(True)
        for col, v in ((2, a), (3, b)):
            c = ws_ass.cell(row=r, column=col, value=datetime.fromisoformat(v))
            c.font, c.fill, c.number_format = f(color=BLUE), INPUT_FILL, DATE
        ws_ass.cell(row=r, column=5, value=desc).font = f()
    ws_ass["A16"] = "Fixed settings"
    ws_ass["A18"], ws_ass["B18"] = "Rolling volatility window (trading days)", ROLL
    ws_ass["E18"] = "Fixed in the Daily calc formulas; about 3 months"
    ws_ass["A19"], ws_ass["B19"] = "Risk-free rate for Sharpe ratio", "BIL"
    ws_ass["E19"] = "1-3 month U.S. T-bill ETF"
    ws_ass["A20"], ws_ass["B20"] = "Rebalancing", "Monthly"
    ws_ass["E20"] = "All portfolios reset to target weights at each month-end"
    for r in (18, 19, 20):
        ws_ass.cell(row=r, column=1).font = f()
        ws_ass.cell(row=r, column=2).font = f()
        ws_ass.cell(row=r, column=5).font = f(color="555555", italic=True)
    TD = "Assumptions!$B$17"

    # ------------------------------------------------------------------ Prices (data)
    title(ws_px, "Prices", f"Daily dividend-adjusted closes from Yahoo Finance. Base day {prices.index[0].date()} is the close before the first return.")
    ws_px.delete_rows(1, 2)
    header(ws_px, 1, ["Date", *ASSETS], widths=[12] + [10] * len(ASSETS))
    for i, (d, row) in enumerate(prices[ASSETS].iterrows()):
        r = 2 + i
        ws_px.cell(row=r, column=1, value=d.to_pydatetime()).number_format = DATE
        for j, t in enumerate(ASSETS):
            c = ws_px.cell(row=r, column=2 + j, value=float(row[t]))
            c.font = Font(name=FONT, color=BLUE)
    ws_px.freeze_panes = "B2"
    pcol = {t: L(2 + j) for j, t in enumerate(ASSETS)}

    # ------------------------------------------------------------------ Daily calc layout
    col = 1
    C = {}
    C["date"] = L(col); col += 1
    C["new"] = L(col); col += 1
    C["ret"] = {t: L(col + j) for j, t in enumerate(ASSETS)}; col += len(ASSETS)
    C["grow"] = {t: L(col + j) for j, t in enumerate(GROWTH)}; col += len(GROWTH)
    C["port"] = {p: L(col + j) for j, p in enumerate(port_names)}; col += len(port_names)
    C["pval"] = {p: L(col + j) for j, p in enumerate(port_names)}; col += len(port_names)
    C["ppeak"] = {p: L(col + j) for j, p in enumerate(port_names)}; col += len(port_names)
    C["pdd"] = {p: L(col + j) for j, p in enumerate(port_names)}; col += len(port_names)
    C["apeak"] = {t: L(col + j) for j, t in enumerate(GROWTH)}; col += len(GROWTH)
    C["add"] = {t: L(col + j) for j, t in enumerate(GROWTH)}; col += len(GROWTH)
    C["vol"] = {t: L(col + j) for j, t in enumerate(GROWTH)}; col += len(GROWTH)

    def rng(c, a=first, b=last):
        return f"'Daily calc'!${c}${a}:${c}${b}"

    DATES = rng(C["date"])
    DATES_ALL = rng(C["date"], 2)

    # ------------------------------------------------------------------ Allocation scenarios (weights live here)
    title(ws_alloc, "Allocation scenarios",
          "60/40 base case vs. adding each strategy. Weights are formulas from Assumptions; statistics are formulas on 'Daily calc'.")
    hdr = ["Portfolio", "Strategy", "Stand-in ETF", "Allocation", "Weight SPY", "Weight AGG", "Weight strategy",
           "Annual return", "Volatility", "Sharpe ratio", "Sharpe vs 60/40", "Max drawdown", "Drawdown vs 60/40",
           "Drawdown peak", "Drawdown trough"]
    header(ws_alloc, 4, hdr, widths=[13, 26, 11, 11, 11, 11, 12, 12, 11, 11, 13, 13, 14, 13, 13])
    AROW = {}
    for i, p in enumerate(port_names):
        r = 5 + i
        AROW[p] = r
        t = None if p == "60/40" else p.split()[-1]
        alloc = "0" if t is None else ("Assumptions!$B$8" if p.startswith("10%") else "Assumptions!$B$9")
        ws_alloc.cell(row=r, column=1, value='="60/40"' if t is None else f'=TEXT({alloc},"0%")&" {t}"')
        ws_alloc.cell(row=r, column=2, value="None (base case)" if t is None else STRATEGIES[t][0])
        ws_alloc.cell(row=r, column=3, value=t or "-")
        ws_alloc.cell(row=r, column=4, value=f"={alloc}" if t else 0)
        ws_alloc.cell(row=r, column=5, value=f"=Assumptions!$B$4*(1-D{r})")
        ws_alloc.cell(row=r, column=6, value=f"=Assumptions!$B$5*(1-D{r})")
        ws_alloc.cell(row=r, column=7, value=f"=D{r}")
        pr, dd, pv, pk = rng(C["port"][p]), rng(C["pdd"][p], 2), rng(C["pval"][p], 2), rng(C["ppeak"][p], 2)
        ws_alloc.cell(row=r, column=8, value=f"=EXP(SUMPRODUCT(LN(1+{pr}))*{TD}/COUNT({pr}))-1")
        ws_alloc.cell(row=r, column=9, value=f"=STDEV({pr})*SQRT({TD})")
        ws_alloc.cell(row=r, column=10, value=f"=(AVERAGE({pr})-AVERAGE({rng(C['ret'][CASH])}))*{TD}/I{r}")
        ws_alloc.cell(row=r, column=11, value=f"=J{r}-$J$5")
        ws_alloc.cell(row=r, column=12, value=f"=MIN({dd})")
        ws_alloc.cell(row=r, column=13, value=f"=(L{r}-$L$5)*100")
        ws_alloc.cell(row=r, column=15, value=f"=INDEX({DATES_ALL},MATCH(L{r},{dd},0))")
        ws_alloc.cell(row=r, column=14, value=f"=INDEX({DATES_ALL},MATCH(INDEX({pk},MATCH(L{r},{dd},0)),{pv},0))")
        fm = [None, None, None, PCT, PCT, PCT, PCT, PCT, PCT, NUM2, '+0.00;-0.00;0.00', PCT, PP, DATE, DATE]
        for j, fmt in enumerate(fm, start=1):
            c = ws_alloc.cell(row=r, column=j)
            c.font = f(bold=(t is None))
            if fmt:
                c.number_format = fmt
            c.border = Border(bottom=THIN)
    note_r = 5 + len(port_names) + 1
    ws_alloc.cell(row=note_r, column=1, value=(
        "Sharpe ratio: average daily return above T-bills (BIL), annualized, divided by annualized volatility. "
        "Max drawdown: largest fall from a previous high in the daily portfolio value. "
        "Drawdown vs 60/40 is in percentage points; positive means a smaller worst loss.")).font = f(color="555555", italic=True)
    ws_alloc.freeze_panes = "B5"

    # ------------------------------------------------------------------ Daily calc formulas
    ws_d.cell(row=1, column=1)
    heads = {C["date"]: "Date", C["new"]: "New month"}
    for t in ASSETS:
        heads[C["ret"][t]] = f"Return {t}"
    for t in GROWTH:
        heads[C["grow"][t]] = f"Growth since month-end {t}"
        heads[C["apeak"][t]] = f"Peak price {t}"
        heads[C["add"][t]] = f"Drawdown {t}"
        heads[C["vol"][t]] = f"Rolling {ROLL}d vol {t}"
    for p in port_names:
        heads[C["port"][p]] = f"Return {p}"
        heads[C["pval"][p]] = f"Value {p}"
        heads[C["ppeak"][p]] = f"Peak {p}"
        heads[C["pdd"][p]] = f"Drawdown {p}"
    for c, h in heads.items():
        cell = ws_d[f"{c}1"]
        cell.value = h
        cell.font = Font(name=FONT, bold=True, color="FFFFFF", size=9)
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
        ws_d.column_dimensions[c].width = 11
    ws_d.row_dimensions[1].height = 42
    ws_d.column_dimensions[C["date"]].width = 12

    strat_growth = {t: C["grow"][t] for t in STRATEGIES}
    A = "'Allocation scenarios'!"
    for i in range(n):
        r = 2 + i
        ws_d[f"{C['date']}{r}"] = f"=Prices!A{r}"
        ws_d[f"{C['date']}{r}"].number_format = DATE
        if r == 2:
            ws_d[f"{C['new']}{r}"] = 0
            for t in GROWTH:
                ws_d[f"{C['grow'][t]}{r}"] = 1
                ws_d[f"{C['apeak'][t]}{r}"] = f"=Prices!{pcol[t]}{r}"
                ws_d[f"{C['add'][t]}{r}"] = 0
            for p in port_names:
                ws_d[f"{C['pval'][p]}{r}"] = 1
                ws_d[f"{C['ppeak'][p]}{r}"] = 1
                ws_d[f"{C['pdd'][p]}{r}"] = 0
            continue
        ws_d[f"{C['new']}{r}"] = f"=IF(MONTH({C['date']}{r})<>MONTH({C['date']}{r-1}),1,0)"
        for t in ASSETS:
            ws_d[f"{C['ret'][t]}{r}"] = f"=Prices!{pcol[t]}{r}/Prices!{pcol[t]}{r-1}-1"
        for t in GROWTH:
            g, rt = C["grow"][t], C["ret"][t]
            ws_d[f"{g}{r}"] = f"=IF({C['new']}{r}=1,1+{rt}{r},{g}{r-1}*(1+{rt}{r}))"
            ws_d[f"{C['apeak'][t]}{r}"] = f"=MAX({C['apeak'][t]}{r-1},Prices!{pcol[t]}{r})"
            ws_d[f"{C['add'][t]}{r}"] = f"=Prices!{pcol[t]}{r}/{C['apeak'][t]}{r}-1"
            if r >= first + ROLL - 1:
                ws_d[f"{C['vol'][t]}{r}"] = f"=STDEV({rt}{r-ROLL+1}:{rt}{r})*SQRT({TD})"
        for p in port_names:
            ar = AROW[p]
            t = "TLT" if p == "60/40" else p.split()[-1]
            gs, ga, gx = C["grow"][EQUITY], C["grow"][BONDS], strat_growth[t]
            ws_, wa, wx = f"{A}$E${ar}", f"{A}$F${ar}", f"{A}$G${ar}"
            num = f"{ws_}*{gs}{r}+{wa}*{ga}{r}+{wx}*{gx}{r}"
            den = f"{ws_}*{gs}{r-1}+{wa}*{ga}{r-1}+{wx}*{gx}{r-1}"
            ws_d[f"{C['port'][p]}{r}"] = f"=({num})/IF({C['new']}{r}=1,1,{den})-1"
            v, pk = C["pval"][p], C["ppeak"][p]
            ws_d[f"{v}{r}"] = f"={v}{r-1}*(1+{C['port'][p]}{r})"
            ws_d[f"{pk}{r}"] = f"=MAX({pk}{r-1},{v}{r})"
            ws_d[f"{C['pdd'][p]}{r}"] = f"={v}{r}/{pk}{r}-1"
    ws_d.freeze_panes = "C2"

    # ------------------------------------------------------------------ Strategy profiles
    title(ws_prof, "Strategy profiles", "Each stand-in on its own, plus the two 60/40 building blocks. All figures are formulas.")
    hdr = ["Strategy", "Stand-in ETF", "What it holds", "Annual return", "Volatility", "Sharpe ratio", "Max drawdown",
           "Drawdown peak", "Drawdown trough", "Correlation with stocks (SPY)", "Correlation with bonds (AGG)",
           "Correlation with 60/40", f"Typical volatility (median {ROLL}d)", f"Highest {ROLL}d volatility",
           "Date of highest", f"Latest {ROLL}d volatility"]
    header(ws_prof, 4, hdr, widths=[25, 10, 40, 11, 11, 10, 12, 12, 12, 13, 13, 12, 13, 12, 12, 12])
    prof_rows = [(t, *STRATEGIES[t]) for t in STRATEGIES] + [
        (EQUITY, "U.S. stocks (S&P 500)", "60/40 building block"), (BONDS, "Core U.S. bonds", "60/40 building block")]
    PROW = {}
    for i, (t, name, desc) in enumerate(prof_rows):
        r = 5 + i
        PROW[t] = r
        rt, dd, pk, vol = rng(C["ret"][t]), rng(C["add"][t], 2), rng(C["apeak"][t], 2), rng(C["vol"][t], first + ROLL - 1)
        price = f"Prices!${pcol[t]}$2:${pcol[t]}${last}"
        vals = [name, t, desc,
                f"=EXP(SUMPRODUCT(LN(1+{rt}))*{TD}/COUNT({rt}))-1",
                f"=STDEV({rt})*SQRT({TD})",
                f"=(AVERAGE({rt})-AVERAGE({rng(C['ret'][CASH])}))*{TD}/E{r}",
                f"=MIN({dd})",
                f"=INDEX({DATES_ALL},MATCH(INDEX({pk},MATCH(G{r},{dd},0)),{price},0))",
                f"=INDEX({DATES_ALL},MATCH(G{r},{dd},0))",
                f"=CORREL({rt},{rng(C['ret'][EQUITY])})",
                f"=CORREL({rt},{rng(C['ret'][BONDS])})",
                f"=CORREL({rt},{rng(C['port']['60/40'])})",
                f"=MEDIAN({vol})", f"=MAX({vol})",
                f"=INDEX({rng(C['date'], first + ROLL - 1)},MATCH(N{r},{vol},0))",
                f"='Daily calc'!{C['vol'][t]}{last}"]
        fm = [None, None, None, PCT, PCT, NUM2, PCT, DATE, DATE, NUM2, NUM2, NUM2, PCT, PCT, DATE, PCT]
        for j, (v, fmt) in enumerate(zip(vals, fm), start=1):
            c = ws_prof.cell(row=r, column=j, value=v)
            c.font = f(bold=(j == 1))
            c.border = Border(bottom=THIN)
            if fmt:
                c.number_format = fmt
            if j == 3:
                c.alignment = Alignment(wrap_text=True, vertical="top")
    ws_prof.freeze_panes = "B5"

    # ------------------------------------------------------------------ Correlations
    title(ws_cor, "Correlations of daily returns", "1 = move together, 0 = unrelated, -1 = move opposite. Full period, then within each stress window.")
    labs = [EQUITY, BONDS, *STRATEGIES, "60/40"]
    colof = {**{t: C["ret"][t] for t in [EQUITY, BONDS, *STRATEGIES]}, "60/40": C["port"]["60/40"]}
    header(ws_cor, 4, [""] + labs, widths=[22] + [10] * len(labs))
    for i, a in enumerate(labs):
        r = 5 + i
        ws_cor.cell(row=r, column=1, value=a).font = f(True)
        for j, b in enumerate(labs):
            c = ws_cor.cell(row=r, column=2 + j, value=f"=CORREL({rng(colof[a])},{rng(colof[b])})")
            c.number_format, c.font = NUM2, f()
    r0 = 5 + len(labs) + 2
    ws_cor.cell(row=r0 - 1, column=1, value="Correlation with stocks (SPY) inside each stress window").font = f(True)
    header(ws_cor, r0, ["Strategy", *STRESS.keys()])
    x = rng(C["ret"][EQUITY])
    for i, t in enumerate(STRATEGIES):
        r = r0 + 1 + i
        ws_cor.cell(row=r, column=1, value=f"{STRATEGIES[t][0]} ({t})").font = f(True)
        y = rng(C["ret"][t])
        for j, k in enumerate(STRESS):
            m = f"({DATES}>Assumptions!$B${13+j})*({DATES}<=Assumptions!$C${13+j})"
            nn, sx, sy = f"SUMPRODUCT({m})", f"SUMPRODUCT({m}*{x})", f"SUMPRODUCT({m}*{y})"
            sxy, sxx, syy = f"SUMPRODUCT({m}*{x}*{y})", f"SUMPRODUCT({m}*{x}*{x})", f"SUMPRODUCT({m}*{y}*{y})"
            formula = f"=({nn}*{sxy}-{sx}*{sy})/SQRT(({nn}*{sxx}-{sx}^2)*({nn}*{syy}-{sy}^2))"
            c = ws_cor.cell(row=r, column=2 + j, value=formula)
            c.number_format, c.font = NUM2, f()
    COR_STRESS_ROW0 = r0 + 1

    # ------------------------------------------------------------------ Stress tests
    title(ws_str, "Stress tests", "Cumulative return over each window (close on start date to close on end date). Windows are set on Assumptions.")
    ks = list(STRESS)
    header(ws_str, 4, ["Portfolio or asset", *[f"{k} return" for k in ks], *[f"{k} vs 60/40" for k in ks]],
           widths=[26] + [12] * (2 * len(ks)))
    items = [(p, C["port"][p], True) for p in port_names] + [(t, C["ret"][t], False) for t in [EQUITY, BONDS, *STRATEGIES]]
    SROW = {}
    for i, (name, c_, is_port) in enumerate(items):
        r = 5 + i + (0 if is_port else 1)
        SROW[name] = r
        lab = f"={A}A{AROW[name]}" if is_port else f"{name} on its own"
        ws_str.cell(row=r, column=1, value=lab).font = f(bold=(name == "60/40"))
        for j, k in enumerate(ks):
            m = f"({DATES}>Assumptions!$B${13+j})*({DATES}<=Assumptions!$C${13+j})"
            cell = ws_str.cell(row=r, column=2 + j, value=f"=EXP(SUMPRODUCT({m}*LN(1+{rng(c_)})))-1")
            cell.number_format, cell.font = PCT, f(bold=(name == "60/40"))
            if is_port:
                cc = ws_str.cell(row=r, column=2 + len(ks) + j, value=f"=({L(2+j)}{r}-{L(2+j)}$5)*100")
                cc.number_format, cc.font = PP, f()
        for j in range(1, 2 + 2 * len(ks)):
            ws_str.cell(row=r, column=j).border = Border(bottom=THIN)
    ws_str.cell(row=5 + len(port_names), column=1, value="Each asset on its own").font = f(True, color="555555")

    # ------------------------------------------------------------------ Summary (mirrors the one-pager)
    title(ws_sum, "Summary", "The figures on the one-pager. Green cells link to the calculation sheets.")
    ws_sum.column_dimensions["A"].width = 30
    for c_ in "BCDEFG":
        ws_sum.column_dimensions[c_].width = 14
    header(ws_sum, 4, ["Portfolio", "Annual return", "Volatility", "Sharpe ratio", "Max drawdown"])
    for i, p in enumerate(port_names):
        r = 5 + i
        ar = AROW[p]
        for j, src in enumerate(["A", "H", "I", "J", "L"]):
            c = ws_sum.cell(row=r, column=1 + j, value=f"={A}{src}{ar}")
            c.font = f(color=GREEN, bold=(p == "60/40"))
            c.number_format = [None, PCT, PCT, NUM2, PCT][j] or "General"
    s0 = 5 + len(port_names) + 2
    header(ws_sum, s0, ["Stress test", *ks])
    for i, p in enumerate(port_names):
        r = s0 + 1 + i
        ws_sum.cell(row=r, column=1, value=f"={A}A{AROW[p]}").font = f(color=GREEN, bold=(p == "60/40"))
        for j in range(len(ks)):
            c = ws_sum.cell(row=r, column=2 + j, value=f"='Stress tests'!{L(2+j)}{SROW[p]}")
            c.font, c.number_format = f(color=GREEN, bold=(p == "60/40")), PCT

    # ------------------------------------------------------------------ Read me
    title(ws_read, "Bond and credit strategies in a 60/40 portfolio: Excel backup")
    m = R["meta"]
    lines = [
        ("What this is", f"The calculations behind the investor one-pager. Daily data from {m['start']} to {m['end']} ({m['days']:,} trading days)."),
        ("Strategy stand-ins", "; ".join(f"{t} = {v[0]}" for t, v in STRATEGIES.items()) + ". Base portfolio: SPY (stocks) and AGG (bonds). T-bills: BIL."),
        ("Data source", "Yahoo Finance daily dividend-adjusted closes, downloaded 2026-10-05; cross-checked against an independent copy for 2016-2026 (identical)."),
        ("How it works", "Prices is the only data. Daily calc turns prices into returns, monthly-rebalanced portfolios, drawdowns and rolling volatility. The other sheets summarize Daily calc with formulas."),
        ("Change a scenario", "Edit the yellow cells on Assumptions (weights, allocation sizes, stress dates). Everything recalculates."),
        ("Color code", "Blue text = data or input. Black = formula. Green = link to another sheet. Yellow fill = input you can change."),
        ("Sheets", "Summary · Allocation scenarios · Stress tests · Strategy profiles · Correlations · Prices · Daily calc"),
        ("Limits", "ETF prices, not fund NAVs: in 2008 some bond ETFs traded below the value of their holdings for days. Past results do not predict future returns. Not investment advice."),
        ("Generated", datetime.now().strftime("%Y-%m-%d %H:%M") + " by make_onepager.py"),
    ]
    ws_read.column_dimensions["A"].width = 20
    ws_read.column_dimensions["B"].width = 110
    for i, (a, b) in enumerate(lines):
        r = 3 + i
        ws_read.cell(row=r, column=1, value=a).font = f(True)
        c = ws_read.cell(row=r, column=2, value=b)
        c.font, c.alignment = f(), Alignment(wrap_text=True, vertical="top")
        ws_read.row_dimensions[r].height = 30
    for ws in wb.worksheets:
        ws.sheet_properties.tabColor = {"Read me": "1F3A5F", "Assumptions": "FFC000", "Summary": "2A78D6"}.get(ws.title, "C9D1DB")
    wb.save(path)
    return {"AROW": AROW, "PROW": PROW, "SROW": SROW, "COR_STRESS_ROW0": COR_STRESS_ROW0, "ks": ks}
